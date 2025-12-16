"""
Excel Exporter for report data.

Uses openpyxl to generate Excel workbooks.
"""

import io
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Optional

from django.core.files.base import ContentFile

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class ExcelExporter:
    """
    Exporter for generating Excel reports.

    Supports headers, data rows, and basic formatting.
    """

    # Style definitions
    HEADER_FILL = PatternFill(
        start_color='4472C4',
        end_color='4472C4',
        fill_type='solid'
    ) if OPENPYXL_AVAILABLE else None

    HEADER_FONT = Font(
        bold=True,
        color='FFFFFF'
    ) if OPENPYXL_AVAILABLE else None

    BORDER = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    ) if OPENPYXL_AVAILABLE else None

    def __init__(self):
        """Initialize Excel exporter."""
        if not OPENPYXL_AVAILABLE:
            raise ImportError(
                "openpyxl is required for Excel export. "
                "Install with: pip install openpyxl"
            )
        self.workbook = Workbook()
        self.workbook.remove(self.workbook.active)  # Remove default sheet

    def add_sheet(
        self,
        name: str,
        headers: list[str],
        data: list[list[Any]],
        column_widths: Optional[list[int]] = None,
    ) -> None:
        """
        Add a sheet with data to the workbook.

        Args:
            name: Sheet name.
            headers: List of column headers.
            data: List of rows (each row is a list of values).
            column_widths: Optional list of column widths.
        """
        ws = self.workbook.create_sheet(title=name[:31])  # Excel limit is 31 chars

        # Write headers
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT
            cell.border = self.BORDER
            cell.alignment = Alignment(horizontal='center', wrap_text=True)

        # Write data
        for row_idx, row_data in enumerate(data, 2):
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=self._format_value(value))
                cell.border = self.BORDER

                # Right-align numbers
                if isinstance(value, (int, float, Decimal)):
                    cell.alignment = Alignment(horizontal='right')

        # Set column widths
        if column_widths:
            for col_idx, width in enumerate(column_widths, 1):
                ws.column_dimensions[get_column_letter(col_idx)].width = width
        else:
            # Auto-fit column widths
            self._auto_fit_columns(ws, headers, data)

        # Freeze header row
        ws.freeze_panes = 'A2'

    def _format_value(self, value: Any) -> Any:
        """Format value for Excel cell."""
        if value is None:
            return ''
        if isinstance(value, Decimal):
            return float(value)
        if isinstance(value, (date, datetime)):
            return value
        return value

    def _auto_fit_columns(
        self,
        ws,
        headers: list[str],
        data: list[list[Any]],
    ) -> None:
        """Auto-fit column widths based on content."""
        for col_idx, header in enumerate(headers, 1):
            max_length = len(str(header))

            for row in data:
                if col_idx <= len(row):
                    cell_value = str(row[col_idx - 1]) if row[col_idx - 1] else ''
                    max_length = max(max_length, len(cell_value))

            # Add some padding and limit max width
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width

    def add_summary_sheet(
        self,
        name: str,
        summary_data: dict[str, Any],
    ) -> None:
        """
        Add a summary sheet with key-value pairs.

        Args:
            name: Sheet name.
            summary_data: Dict of label-value pairs.
        """
        ws = self.workbook.create_sheet(title=name[:31])

        row = 1
        for label, value in summary_data.items():
            # Label cell
            label_cell = ws.cell(row=row, column=1, value=label)
            label_cell.font = Font(bold=True)
            label_cell.border = self.BORDER

            # Value cell
            value_cell = ws.cell(row=row, column=2, value=self._format_value(value))
            value_cell.border = self.BORDER

            row += 1

        # Set column widths
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 20

    def get_file_content(self) -> bytes:
        """
        Get the workbook content as bytes.

        Returns:
            Bytes content of the Excel file.
        """
        buffer = io.BytesIO()
        self.workbook.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    def save_to_file(self, filename: str) -> None:
        """
        Save the workbook to a file.

        Args:
            filename: Path to save the file.
        """
        self.workbook.save(filename)

    def get_content_file(self, filename: str) -> ContentFile:
        """
        Get the workbook as a Django ContentFile.

        Args:
            filename: Name for the file.

        Returns:
            ContentFile instance.
        """
        content = self.get_file_content()
        return ContentFile(content, name=filename)


def export_report_to_excel(
    report_data: dict,
    report_name: str,
) -> bytes:
    """
    Convenience function to export report data to Excel.

    Args:
        report_data: Dict with 'headers', 'data', and optional 'summary'.
        report_name: Name for the report sheet.

    Returns:
        Bytes content of the Excel file.
    """
    if not OPENPYXL_AVAILABLE:
        raise ImportError("openpyxl is required for Excel export")

    exporter = ExcelExporter()

    # Add summary if present
    if 'summary' in report_data:
        exporter.add_summary_sheet('Summary', report_data['summary'])

    # Add main data sheet
    headers = report_data.get('headers', [])
    data = report_data.get('data', [])
    exporter.add_sheet(report_name, headers, data)

    return exporter.get_file_content()
